#!/usr/bin/env python3
"""
manifest_to_excel.py

Convert indices and data-source JSON manifests into a multi-sheet Excel workbook.

Usage:
  python pipeline/2_index_exploration/index_definition_to_excel.py --manifest pipeline/1_index_construction/indices_manifest.json --out pipeline/1_index_construction/indices_manifest.xlsx

Requires:
  pip install openpyxl pandas
"""

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional

import pandas as pd

try:
    import rasterio
except Exception:  # pragma: no cover - optional dependency for metadata enrichment
    rasterio = None

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.append(str(REPO_ROOT))

from pipeline.utils.paths import load_paths, repo_data_path  # noqa: E402
from pipeline.utils.io import read_csv_robust


def _as_str(x: Any) -> str:
    """Safe stringifier for Excel cells."""
    if x is None:
        return ""
    if isinstance(x, (dict, list)):
        # Keep readable but compact
        return json.dumps(x, ensure_ascii=False)
    return str(x)


def _list_tifs(folder_path: Path) -> List[Path]:
    tifs: List[Path] = []
    for pattern in ("*.tif", "*.tiff", "*.TIF", "*.TIFF"):
        tifs.extend(folder_path.glob(pattern))
    return sorted(tifs)


def _format_resolution(rx: Any, ry: Any) -> str:
    if rx in ("", None) or ry in ("", None):
        return ""
    try:
        return f"{float(abs(rx)):.6g} x {float(abs(ry)):.6g}"
    except Exception:
        return ""


def _format_years(years: set[int]) -> str:
    if not years:
        return ""
    sorted_years = sorted(years)
    ranges = []
    start = prev = sorted_years[0]
    for year in sorted_years[1:]:
        if year == prev + 1:
            prev = year
            continue
        ranges.append(str(start) if start == prev else f"{start}-{prev}")
        start = prev = year
    ranges.append(str(start) if start == prev else f"{start}-{prev}")
    return ", ".join(ranges)


def _format_year_range(years: set[int]) -> str:
    if not years:
        return ""
    years_sorted = sorted(years)
    if len(years_sorted) == 1:
        return str(years_sorted[0])
    return f"{years_sorted[0]}-{years_sorted[-1]}"


def _infer_years_from_filenames(file_names: List[str]) -> Dict[str, str]:
    """
    Infer observation/source years from known raster filename conventions.

    This is intentionally conservative:
    - GHSL-style E2010/E2018/E2022 tokens are treated as source years.
    - 2010_2012 or 201204-201212 tokens are treated as source ranges.
    - R2025A/c202205... tokens are recorded as release/creation years, not
      source years.
    - AGE filenames such as 1975052020 are interpreted as 1975-2020.
    """
    source_years: set[int] = set()
    release_years: set[int] = set()
    notes: set[str] = set()

    for name in file_names:
        lower = name.lower()

        for match in re.finditer(r"E((?:19|20)\d{2})(?!\d)", name):
            source_years.add(int(match.group(1)))

        for match in re.finditer(r"(?<!\d)((?:19|20)\d{2})[_-]((?:19|20)\d{2})(?!\d)", name):
            start, end = int(match.group(1)), int(match.group(2))
            source_years.update(range(min(start, end), max(start, end) + 1))

        for match in re.finditer(r"(?<!\d)((?:19|20)\d{2})\d{2}[-_]?((?:19|20)\d{2})\d{2}(?!\d)", name):
            start, end = int(match.group(1)), int(match.group(2))
            source_years.update(range(min(start, end), max(start, end) + 1))

        age_match = re.search(r"age_((?:19|20)\d{2})\d{2}((?:19|20)\d{2})", lower)
        if age_match:
            start, end = int(age_match.group(1)), int(age_match.group(2))
            source_years.update(range(min(start, end), max(start, end) + 1))
            notes.add("AGE filename interpreted as multi-epoch source range")

        for match in re.finditer(r"R((?:19|20)\d{2})[A-Z]?", name):
            release_years.add(int(match.group(1)))
        for match in re.finditer(r"c((?:19|20)\d{2})\d{4,}", name, flags=re.IGNORECASE):
            release_years.add(int(match.group(1)))

    if not source_years:
        notes.add("no source year inferred from TIFF filenames")

    return {
        "source_years": _format_years(source_years),
        "source_year_range": _format_year_range(source_years),
        "release_or_creation_years": _format_years(release_years),
        "year_inference_notes": "; ".join(sorted(notes)),
    }


def _inspect_raster_folder(folder_name: str, folder_path: Optional[Path]) -> Dict[str, Any]:
    """
    Inspect a curated raster folder without changing the JSON manifest.

    The index runner uses the coarsest input grid as the reference grid per
    index. Therefore the most useful folder-level value here is the coarsest
    native pixel size among TIFFs in the folder.
    """
    row: Dict[str, Any] = {
        "folder_path": _as_str(folder_path) if folder_path is not None else "",
        "tif_count": "",
        "crs": "",
        "resolution_x": "",
        "resolution_y": "",
        "resolution": "",
        "resolution_unit": "",
        "pixel_area": "",
        "unique_resolutions": "",
        "source_years": "",
        "source_year_range": "",
        "release_or_creation_years": "",
        "sample_tif": "",
        "resolution_notes": "",
        "year_inference_notes": "",
    }

    if folder_path is None:
        row["resolution_notes"] = "not inspected: data_root path could not be resolved"
        return row
    if rasterio is None:
        row["resolution_notes"] = "not inspected: rasterio is not available"
        return row
    if not folder_path.exists():
        row["resolution_notes"] = "not inspected: folder does not exist"
        return row

    tifs = _list_tifs(folder_path)
    row["tif_count"] = len(tifs)
    if not tifs:
        row["resolution_notes"] = "not inspected: no TIFF files found"
        return row

    row.update(_infer_years_from_filenames([tif.name for tif in tifs]))

    best: Optional[Dict[str, Any]] = None
    unique_resolutions = set()
    crs_values = set()
    errors = []
    for tif in tifs:
        try:
            with rasterio.open(tif) as src:
                rx, ry = src.res
                rx_abs = float(abs(rx))
                ry_abs = float(abs(ry))
                area = rx_abs * ry_abs
                unique_resolutions.add((round(rx_abs, 8), round(ry_abs, 8)))
                crs_text = src.crs.to_string() if src.crs is not None else ""
                if crs_text:
                    crs_values.add(crs_text)
                candidate = {
                    "resolution_x": rx_abs,
                    "resolution_y": ry_abs,
                    "pixel_area": area,
                    "crs": crs_text,
                    "sample_tif": tif.name,
                }
                if best is None or area > best["pixel_area"]:
                    best = candidate
        except Exception as exc:
            errors.append(f"{tif.name}: {exc}")

    if best is None:
        row["resolution_notes"] = "not inspected: could not open TIFF files"
        return row

    row.update(best)
    row["resolution"] = _format_resolution(best["resolution_x"], best["resolution_y"])
    row["resolution_unit"] = "CRS units; meters for Mollweide"
    row["unique_resolutions"] = "; ".join(
        _format_resolution(rx, ry) for rx, ry in sorted(unique_resolutions)
    )
    notes = []
    if len(unique_resolutions) > 1:
        notes.append("mixed native resolutions; coarsest resolution reported")
    if len(crs_values) > 1:
        notes.append("mixed CRS values")
    if errors:
        notes.append(f"{len(errors)} TIFF(s) could not be inspected")
    row["resolution_notes"] = "; ".join(notes)
    return row


def _build_folder_resolution_catalog(
    manifest: Dict[str, Any],
    paths_cfg: Optional[Dict[str, Any]],
) -> Dict[str, Dict[str, Any]]:
    data_root = paths_cfg["sat_data_curated"] if paths_cfg else None
    folder_catalog = manifest.get("folder_catalog", {})
    if not isinstance(folder_catalog, dict):
        return {}

    return {
        folder_name: _inspect_raster_folder(
            folder_name,
            data_root / folder_name if data_root is not None else None,
        )
        for folder_name in folder_catalog
    }


def _index_resolution_summary(
    inputs: List[Any],
    folder_resolution_catalog: Dict[str, Dict[str, Any]],
) -> Dict[str, Any]:
    rows = []
    for folder in inputs:
        meta = folder_resolution_catalog.get(str(folder), {})
        rx = meta.get("resolution_x", "")
        ry = meta.get("resolution_y", "")
        area = meta.get("pixel_area", "")
        rows.append({
            "folder": str(folder),
            "resolution": _format_resolution(rx, ry),
            "pixel_area": area,
        })

    available = [
        row for row in rows
        if row.get("pixel_area") not in ("", None) and pd.notna(row.get("pixel_area"))
    ]
    ref = max(available, key=lambda row: float(row["pixel_area"])) if available else {}
    return {
        "input_grid_resolutions": "; ".join(
            f"{row['folder']}={row['resolution']}" for row in rows if row["resolution"]
        ),
        "reference_grid_folder": ref.get("folder", ""),
        "reference_grid_resolution": ref.get("resolution", ""),
        "reference_grid_pixel_area": ref.get("pixel_area", ""),
    }


def _index_year_summary(
    inputs: List[Any],
    folder_resolution_catalog: Dict[str, Dict[str, Any]],
) -> Dict[str, Any]:
    rows = []
    all_years: set[int] = set()
    for folder in inputs:
        meta = folder_resolution_catalog.get(str(folder), {})
        years_text = str(meta.get("source_years", "") or "")
        if years_text:
            for part in years_text.split(","):
                part = part.strip()
                if "-" in part:
                    start, end = part.split("-", 1)
                    if start.strip().isdigit() and end.strip().isdigit():
                        all_years.update(range(int(start), int(end) + 1))
                elif part.isdigit():
                    all_years.add(int(part))
            rows.append(f"{folder}={years_text}")
    return {
        "input_source_years": "; ".join(rows),
        "input_source_year_range": _format_year_range(all_years),
    }


def _build_admin_indicator_table() -> pd.DataFrame:
    admin_csv = repo_data_path("all_used_adm_indicators.csv")
    split_json = repo_data_path("adm_data_split.json")

    split: Dict[str, List[str]] = {}
    if split_json is not None and split_json.exists():
        with split_json.open("r", encoding="utf-8") as f:
            raw_split = json.load(f)
        if isinstance(raw_split, dict):
            split = {
                str(group): [str(col) for col in cols]
                for group, cols in raw_split.items()
                if isinstance(cols, list)
            }

    csv_columns: List[str] = []
    row_count = ""
    if admin_csv is not None and admin_csv.exists():
        admin_df = read_csv_robust(admin_csv)
        csv_columns = [str(col) for col in admin_df.columns]
        row_count = str(len(admin_df)) if "ID" in csv_columns else ""

    category_by_indicator: Dict[str, str] = {}
    for category, indicators in split.items():
        label = "basic_prediction" if category == "basic" else "complex_detailed"
        for indicator in indicators:
            category_by_indicator[indicator] = label

    indicators = sorted(set(category_by_indicator) | {col for col in csv_columns if col != "ID"})
    rows = []
    for indicator in indicators:
        category = category_by_indicator.get(indicator, "not_classified_in_split")
        rows.append({
            "indicator": indicator,
            "category": category,
            "used_for_prediction": category == "basic_prediction",
            "reserved_for_residual_analysis": category == "complex_detailed",
            "available_in_admin_csv": indicator in csv_columns,
            "admin_csv": _as_str(admin_csv),
            "admin_split_json": _as_str(split_json),
            "admin_csv_rows": row_count,
            "notes": "" if indicator in csv_columns else "listed in split file but not present in admin CSV",
        })
    return pd.DataFrame(rows)


def load_manifest(path: Path) -> Dict[str, Any]:
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def build_source_tables(registry: Dict[str, Any], manifest: Dict[str, Any]):
    """Validate source identifiers and export provenance without promoting planned claims."""
    sources = registry["sources"]
    ids = [source["id"] for source in sources]
    if len(set(ids)) != len(ids):
        raise ValueError("Duplicate data source IDs")
    evidence_ids = {item["id"] for item in registry.get("evidence", [])}
    folder_sources: Dict[str, List[str]] = {}
    admin_sources: Dict[str, List[str]] = {}
    for source in sources:
        if not isinstance(source.get("verified"), bool):
            raise ValueError(f"Source {source['id']} must have a boolean verified field")
        if set(source.get("evidence_ids", [])) - evidence_ids:
            raise ValueError(f"Source {source['id']} references unknown evidence")
        folders = source.get("raster_folders", [])
        if set(folders) - set(manifest.get("folder_catalog", {})):
            raise ValueError(f"Source {source['id']} references unknown raster folders")
        if source.get("usage_status") == "planned_only" and (folders or source.get("administrative_indicators")):
            raise ValueError(f"Planned source {source['id']} cannot be mapped to pipeline inputs")
        for folder in folders:
            folder_sources.setdefault(folder, []).append(source["id"])
        for indicator in source.get("administrative_indicators", []):
            admin_sources.setdefault(indicator, []).append(source["id"])

    # Lists become readable multiline cells; null remains an empty Excel cell.
    def cell(value):
        if isinstance(value, list) and all(isinstance(v, str) for v in value):
            return "\n".join(value)
        return _as_str(value) if isinstance(value, (dict, list)) else value

    tables = {"Data_Sources": pd.DataFrame([
        {key: cell(value) for key, value in source.items()} for source in sources
    ])}
    evidence_rows = []
    for evidence in registry.get("evidence", []):
        for page in evidence.get("pages", [{"page": None, "text": evidence.get("text")} ]):
            evidence_rows.append({
                "evidence_id": evidence["id"], "path": evidence.get("path"),
                "verified": evidence.get("verified"), "page": page["page"], "text": page["text"],
            })
    tables["Source_Evidence"] = pd.DataFrame(evidence_rows)
    tables["Source_Field_Guide"] = pd.DataFrame([
        {"field": key, "meaning": value} for key, value in registry.get("field_notes", {}).items()
    ])
    return tables, folder_sources, admin_sources


def build_tables(
    manifest: Dict[str, Any],
    paths_cfg: Optional[Dict[str, Any]] = None,
    sources_manifest: Optional[Dict[str, Any]] = None,
) -> Dict[str, pd.DataFrame]:
    if sources_manifest is None:
        sources_manifest = load_manifest(REPO_ROOT / "data" / "data_sources_manifest.json")
    tables, folder_sources, admin_sources = build_source_tables(sources_manifest, manifest)

    # --- Project / top-level info ---
    project = manifest.get("project", {})
    defaults = manifest.get("global_defaults", {})
    helpers = manifest.get("helpers", {})
    folder_resolution_catalog = _build_folder_resolution_catalog(manifest, paths_cfg)

    project_rows = []
    for section_name, section in [
        ("project", project),
        ("global_defaults", defaults),
        ("helpers", helpers),
    ]:
        if isinstance(section, dict):
            for k, v in section.items():
                project_rows.append(
                    {"section": section_name, "key": k, "value": _as_str(v)}
                )
        else:
            project_rows.append({"section": section_name, "key": "", "value": _as_str(section)})

    tables["Project"] = pd.DataFrame(project_rows)
    tables["Paths_Config"] = pd.DataFrame(
        [
            {"key": str(key), "path": _as_str(value)}
            for key, value in sorted((paths_cfg or {}).items())
        ]
    )

    # --- Folder catalog ---
    folder_catalog = manifest.get("folder_catalog", {})
    folder_rows = []
    if isinstance(folder_catalog, dict):
        for folder_name, meta in folder_catalog.items():
            row = {"folder": folder_name, "source_ids": "; ".join(folder_sources.get(folder_name, []))}
            if isinstance(meta, dict):
                for k, v in meta.items():
                    row[k] = _as_str(v)
            else:
                row["meta"] = _as_str(meta)
            for k, v in folder_resolution_catalog.get(folder_name, {}).items():
                row[k] = _as_str(v)
            folder_rows.append(row)
    tables["FolderCatalog"] = pd.DataFrame(folder_rows)

    # --- Indices: summary + nested tables ---
    indices = manifest.get("indices", [])
    if not isinstance(indices, list):
        raise ValueError("manifest['indices'] must be a list")

    summary_rows: List[Dict[str, Any]] = []
    cell_vars_rows: List[Dict[str, Any]] = []
    cell_filters_rows: List[Dict[str, Any]] = []
    reduce_rows: List[Dict[str, Any]] = []
    reduce_stats_rows: List[Dict[str, Any]] = []
    outputs_rows: List[Dict[str, Any]] = []

    for idx in indices:
        if not isinstance(idx, dict):
            continue

        idx_id = idx.get("id", "")
        group = idx.get("group", "")
        desc = idx.get("description", "")
        inputs = idx.get("inputs", [])
        bind = idx.get("bind", {})

        cell = idx.get("cell", {}) or {}
        reduce = idx.get("reduce", {}) or {}
        outputs = idx.get("outputs", []) or []

        resolution_summary = _index_resolution_summary(
            inputs if isinstance(inputs, list) else [],
            folder_resolution_catalog,
        )
        year_summary = _index_year_summary(
            inputs if isinstance(inputs, list) else [],
            folder_resolution_catalog,
        )

        # Summary row
        summary_rows.append(
            {
                "id": idx_id,
                "group": group,
                "description": desc,
                "inputs": _as_str(inputs),
                "source_ids": "; ".join(sorted({
                    source_id for folder in inputs for source_id in folder_sources.get(folder, [])
                })),
                "input_grid_resolutions": _as_str(resolution_summary["input_grid_resolutions"]),
                "reference_grid_folder": _as_str(resolution_summary["reference_grid_folder"]),
                "reference_grid_resolution": _as_str(resolution_summary["reference_grid_resolution"]),
                "reference_grid_pixel_area": _as_str(resolution_summary["reference_grid_pixel_area"]),
                "input_source_years": _as_str(year_summary["input_source_years"]),
                "input_source_year_range": _as_str(year_summary["input_source_year_range"]),
                "bind": _as_str(bind),
                "cell_value": _as_str(cell.get("value")),
                "cell_mask": _as_str(cell.get("mask")),
                "reduce_method": _as_str(reduce.get("method")),
                "reduce_weight": _as_str(reduce.get("weight")),
                "reduce_numerator": _as_str(reduce.get("numerator")),
                "reduce_denominator": _as_str(reduce.get("denominator")),
                "reduce_condition": _as_str(reduce.get("condition")),
                "notes": _as_str(idx.get("notes")),
            }
        )

        # Cell vars
        vars_dict = cell.get("vars", {}) or {}
        if isinstance(vars_dict, dict):
            for var_name, expr in vars_dict.items():
                cell_vars_rows.append(
                    {
                        "index_id": idx_id,
                        "var_name": var_name,
                        "expression": _as_str(expr),
                    }
                )

        # Cell filters
        filters = cell.get("filters", []) or []
        if isinstance(filters, list):
            for f in filters:
                if isinstance(f, dict):
                    row = {"index_id": idx_id}
                    # keep stable columns even if some absent
                    for k in ["type", "var", "q", "op", "value", "notes"]:
                        if k in f:
                            row[k] = _as_str(f.get(k))
                    # include any other keys too
                    for k, v in f.items():
                        if k not in row:
                            row[k] = _as_str(v)
                    cell_filters_rows.append(row)
                else:
                    cell_filters_rows.append({"index_id": idx_id, "filter": _as_str(f)})

        # Reduce table
        reduce_row = {"index_id": idx_id}
        if isinstance(reduce, dict):
            for k, v in reduce.items():
                if k == "stats":
                    continue
                reduce_row[k] = _as_str(v)
        reduce_rows.append(reduce_row)

        # Reduce stats
        stats = reduce.get("stats", []) if isinstance(reduce, dict) else []
        if isinstance(stats, list):
            for s in stats:
                if isinstance(s, dict):
                    row = {"index_id": idx_id}
                    for k, v in s.items():
                        row[k] = _as_str(v)
                    reduce_stats_rows.append(row)
                else:
                    reduce_stats_rows.append({"index_id": idx_id, "stat": _as_str(s)})

        # Outputs
        if isinstance(outputs, list):
            for out in outputs:
                if isinstance(out, dict):
                    row = {"index_id": idx_id}
                    for k, v in out.items():
                        row[k] = _as_str(v)
                    outputs_rows.append(row)
                else:
                    outputs_rows.append({"index_id": idx_id, "output": _as_str(out)})

    tables["Indices_Summary"] = pd.DataFrame(summary_rows)
    tables["Cell_Vars"] = pd.DataFrame(cell_vars_rows)
    tables["Cell_Filters"] = pd.DataFrame(cell_filters_rows)
    tables["Reduce"] = pd.DataFrame(reduce_rows)
    tables["Reduce_Stats"] = pd.DataFrame(reduce_stats_rows)
    tables["Outputs"] = pd.DataFrame(outputs_rows)
    tables["Administrative_Indicators"] = _build_admin_indicator_table()
    admin_table = tables["Administrative_Indicators"]
    if "indicator" in admin_table:
        admin_table.insert(1, "source_ids", admin_table["indicator"].map(
            lambda indicator: "; ".join(admin_sources.get(indicator, []))
        ))

    coverage_rows = [
        {"input_type": "raster_folder", "input": folder,
         "source_ids": "; ".join(folder_sources.get(folder, [])),
         "mapping_status": "mapped" if folder_sources.get(folder) else "unmapped"}
        for folder in manifest.get("folder_catalog", {})
    ]
    if "indicator" in admin_table:
        coverage_rows.extend(
            {"input_type": "administrative_indicator", "input": indicator,
             "source_ids": "; ".join(admin_sources.get(indicator, [])),
             "mapping_status": "mapped" if admin_sources.get(indicator) else "unmapped"}
            for indicator in admin_table["indicator"]
        )
    tables["Source_Coverage"] = pd.DataFrame(coverage_rows)

    # --- README sheet content as a small DF (easy write) ---
    readme_lines = [
        "This workbook is generated from indices and data-source JSON manifests.",
        "Unknown source metadata is blank. Verified applies only within verification_scope; citations are assessed separately.",
        "planned_only sources are a priori candidates, not confirmed pipeline inputs. See notes for differences from downloaded data.",
        "",
        "Sheets:",
        "- Data_Sources: source overview, links, download notes, candidate APA citations, verification and pipeline mappings",
        "- Source_Evidence: supplied download notes and unverified a priori PDF text (by page)",
        "- Source_Field_Guide: registry field semantics",
        "- Source_Coverage: mapped and unmapped raster folders and administrative indicators",
        "- Project: top-level metadata, defaults, helpers",
        "- Paths_Config: filesystem paths loaded from pipeline/config/paths.json",
        "- FolderCatalog: available raster folders (logical sources)",
        "- Indices_Summary: one row per index (high-level overview, including input/reference grid resolutions when raster folders can be inspected)",
        "- Cell_Vars: one row per computed cell variable",
        "- Cell_Filters: one row per filter condition",
        "- Reduce: reduction/aggregation configuration per index",
        "- Reduce_Stats: additional distribution stats per index",
        "- Outputs: final output columns produced per index",
        "- Administrative_Indicators: administrative indicators split into basic prediction variables and complex/detailed residual-analysis variables",
        "",
        "Tip: You can filter Indices_Summary by group to quickly review definitions.",
    ]
    tables["README"] = pd.DataFrame({"README": readme_lines})

    return tables


def autosize_columns(writer, sheet_name: str, df: pd.DataFrame, max_width: int = 60):
    """Autosize columns in openpyxl worksheet (bounded)."""
    ws = writer.sheets[sheet_name]
    for col_idx, col_name in enumerate(df.columns, start=1):
        # compute best width based on header + a sample of rows
        series = df[col_name].astype(str)
        # avoid huge cost: look at first N rows only
        N = min(len(series), 500)
        max_len = max([len(str(col_name))] + [len(str(series.iloc[i])) for i in range(N)]) if N > 0 else len(str(col_name))
        width = min(max_len + 2, max_width)
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width


def write_excel(tables: Dict[str, pd.DataFrame], out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)

    # Use openpyxl engine for formatting
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        # Order matters (README first)
        sheet_order = [
            "README",
            "Data_Sources",
            "Source_Coverage",
            "Source_Field_Guide",
            "Source_Evidence",
            "Project",
            "Paths_Config",
            "FolderCatalog",
            "Indices_Summary",
            "Cell_Vars",
            "Cell_Filters",
            "Reduce",
            "Reduce_Stats",
            "Outputs",
            "Administrative_Indicators",
        ]

        # Write all tables, even if empty
        for name in sheet_order:
            df = tables.get(name, pd.DataFrame())
            df.to_excel(writer, sheet_name=name, index=False)
            autosize_columns(writer, name, df)

        # Write any additional tables not in sheet_order
        for name, df in tables.items():
            if name in sheet_order:
                continue
            df.to_excel(writer, sheet_name=name[:31], index=False)
            autosize_columns(writer, name[:31], df)

        from openpyxl.styles import Alignment, Font, PatternFill

        for name, ws in writer.sheets.items():
            ws.freeze_panes = "B2" if name == "Data_Sources" else "A2"
            ws.sheet_view.showGridLines = False
            if ws.max_column and name != "README":
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="244763")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[1].height = 30
            headers = {cell.value: cell.column for cell in ws[1]}
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    # Export source text literally, including text beginning with '='.
                    if cell.data_type == "f":
                        cell.data_type = "s"
                    if isinstance(cell.value, str) and cell.value.startswith(("https://", "http://")) and not any(c.isspace() for c in cell.value):
                        cell.hyperlink = cell.value
                        cell.font = Font(color="0563C1", underline="single")
                if name == "Data_Sources":
                    verified = row[headers["verified"] - 1].value
                    row[headers["verified"] - 1].fill = PatternFill("solid", fgColor="E2F0D9" if verified else "FFF2CC")
                    ws.row_dimensions[row[0].row].height = 90
            if name == "README":
                ws.column_dimensions["A"].width = 115
                for row in range(2, ws.max_row + 1):
                    ws.row_dimensions[row].height = 32
            if name == "Source_Evidence":
                ws.column_dimensions["E"].width = 110
                for row in range(2, ws.max_row + 1):
                    ws.row_dimensions[row].height = 120

    print(f"Wrote Excel: {out_path}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--manifest", required=True, help="Path to indices_manifest.json")
    ap.add_argument("--out", required=True, help="Output .xlsx path")
    ap.add_argument(
        "--sources-manifest", type=Path,
        default=REPO_ROOT / "data" / "data_sources_manifest.json",
        help="Data-source registry JSON (defaults to the repo data/data_sources_manifest.json)",
    )
    ap.add_argument(
        "--paths-config",
        default=None,
        help=(
            "Optional path to pipeline path config. Defaults to pipeline/config/paths.json."
        ),
    )
    args = ap.parse_args()

    manifest_path = Path(args.manifest)
    out_path = Path(args.out)

    manifest = load_manifest(manifest_path)
    PATHS = load_paths(args.paths_config)
    tables = build_tables(manifest, paths_cfg=PATHS, sources_manifest=load_manifest(args.sources_manifest))
    write_excel(tables, out_path)


if __name__ == "__main__":
    main()
